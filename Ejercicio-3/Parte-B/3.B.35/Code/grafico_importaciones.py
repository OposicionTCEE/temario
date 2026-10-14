#!/usr/bin/env python3
"""Genera Img/USA_importaciones.tikz: importaciones de mercancías de EE. UU. desde México, Canadá y China.

(a) en miles de millones de dólares corrientes; (b) en % de las importaciones totales de EE. UU.
Lee la hoja USA_importaciones de Code/datos_3B35.xlsx (Code/datos_comercio.py USA importaciones 1994 2024 MEX,CAN,CHN).
Uso, desde la carpeta del tema (con openpyxl):
  /tmp/v/bin/python Code/grafico_importaciones.py && bash Code/generar_figura.sh USA_importaciones
"""
from pathlib import Path

SERIES = [('MEX', 'México', (60, 140, 90)), ('CAN', 'Canadá', (200, 60, 60)), ('CHN', 'China', (233, 160, 40))]
HITOS = [(1994, 'NAFTA'), (2001, 'China, OMC'), (2018, 'Aranceles a China'), (2020, 'USMCA')]


def main():
    from openpyxl import load_workbook
    tema = Path(__file__).resolve().parent.parent
    ws = load_workbook(tema / 'Code' / 'datos_3B35.xlsx')['USA_importaciones']
    cab = [c for c in next(ws.iter_rows(min_row=4, max_row=4, values_only=True))]
    filas = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    col = {s: cab.index(s) for s, _, _ in SERIES}
    itot = cab.index('Total (WLD)')
    o = ['% Generado por Code/grafico_importaciones.py con la hoja USA_importaciones.']
    for i, (_, _, rgb) in enumerate(SERIES):
        o.append('\\definecolor{ci%d}{RGB}{%d,%d,%d}' % (i, *rgb))
    o.append('\\begin{tikzpicture}[font=\\sffamily]')
    for p, (titulo, ylab, f) in enumerate([
            ('(a) Miles de millones de USD', 'miles de millones de USD', lambda r, c: r[c] / 1000),
            ('(b) \\% de las importaciones de EE.\\,UU.', '\\% del total', lambda r, c: 100 * r[c] / r[itot])]):
        pos = 'name=a' if p == 0 else 'at={(a.east)}, anchor=west, xshift=1.2cm'
        ymax = 620 if p == 0 else 25
        o += [f'  \\begin{{axis}}[{pos}, width=7.9cm, height=6cm, xmin=1993, xmax=2024, ymin=0, ymax={ymax},',
              '      xtick={1995,2000,2005,2010,2015,2020}, x tick label style={/pgf/number format/1000 sep={}},',
              f'      ylabel={{{ylab}}}, ylabel style={{font=\\sffamily\\scriptsize}}, tick label style={{font=\\sffamily\\scriptsize}},',
              f'      title={{{titulo}}}, title style={{font=\\sffamily\\small}}, axis lines*=left, ymajorgrids, grid style={{black!8}},',
              '      legend style={font=\\sffamily\\scriptsize, draw=none, fill=none, at={(0.02,0.98)}, anchor=north west}, legend cell align=left]']
        for i, (s, nom, _) in enumerate(SERIES):
            coords = ' '.join(f'({r[0]},{f(r, col[s]):.2f})' for r in filas)
            o.append(f'    \\addplot[ci{i}, line width=1.2pt, mark=none] coordinates {{{coords}}};')
            if p == 0:
                o.append(f'    \\addlegendentry{{{nom}}}')
        if p == 1:
            for anio, txt in HITOS:
                o.append(f'    \\draw[black!40, densely dashed, line width=0.4pt] (axis cs:{anio},0) -- (axis cs:{anio},{ymax});')
                o.append(f'    \\node[rotate=90, anchor=south west, font=\\sffamily\\tiny, text=black!70, inner sep=1.2pt] '
                         f'at (axis cs:{anio},0.4) {{{txt}}};')
        o.append('  \\end{axis}')
    o.append('\\end{tikzpicture}')
    (tema / 'Img' / 'USA_importaciones.tikz').write_text('\n'.join(o) + '\n', encoding='utf-8')
    print('Escrito Img/USA_importaciones.tikz')


if __name__ == '__main__':
    main()
