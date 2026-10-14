#!/usr/bin/env python3
"""Genera Img/AfCFTA_comercio_intrarregional.tikz (TikZ puro, dos paneles):
(a) cuota del comercio intraafricano en las exportaciones de los países africanos que declaran, 1996-2023;
(b) cuota del comercio intrarregional en 2023 en los bloques del tema y en la UE-27 y la ASEAN, como referencia.

Lee las hojas <GRUPO>_intrarregional de Code/datos_3B35.xlsx (las escribe Code/datos_comercio.py). La EFTA no tiene hoja:
su cuota se calcula con Code/datos/destinos.csv (0,46 % en 2023).
Uso, desde la carpeta del tema (con openpyxl):
  /tmp/v/bin/python Code/grafico_africa.py
  bash Code/generar_figura.sh AfCFTA_comercio_intrarregional
"""
from pathlib import Path

TEMA = Path(__file__).resolve().parent.parent
BLOQUES = [('UE27', 'UE-27'), ('NAFTA', 'NAFTA/USMCA'), ('MCCA', 'MCCA'), ('ASEAN', 'ASEAN'), ('CEFTA', 'CEFTA'),
           ('AFCFTA', 'AfCFTA'), ('MERCOSUR', 'MERCOSUR'), ('CAN', 'CAN'), ('EFTA', 'EFTA')]
EFTA_2023 = 0.46
HITOS = [(2002, 'Unión Africana'), (2012, 'decisión'), (2018, 'firma'), (2021, 'comercio')]


def serie(ws):
    return {r[0]: (r[3], r[4]) for r in ws.iter_rows(min_row=6, values_only=True) if isinstance(r[0], int)}


def main():
    from openpyxl import load_workbook
    wb = load_workbook(TEMA / 'Code' / 'datos_3B35.xlsx', read_only=True)
    afr = serie(wb['AFCFTA_intrarregional'])
    puntos = [(a, q) for a, (q, n) in sorted(afr.items()) if q is not None and 1996 <= a <= 2023]
    o = ['% Generado por Code/grafico_africa.py.', '\\definecolor{cafr}{RGB}{60,130,90}', '\\definecolor{cotro}{RGB}{150,150,150}',
         '\\begin{tikzpicture}[font=\\sffamily]']
    # (a) Serie
    W, H, a0, a1, ymax = 7.2, 4.2, 1995, 2024, 25

    def px(a):
        return (a - a0) / (a1 - a0) * W

    def py(q):
        return q / ymax * H
    o.append(f'  \\node[anchor=west, font=\\sffamily\\small] at (-0.6,{H + 0.75:.2f}) {{(a) Comercio intraafricano}};')
    for q in range(0, ymax + 1, 5):
        o.append(f'  \\draw[black!12, line width=0.3pt] (0,{py(q):.2f}) -- ({W:.2f},{py(q):.2f}); '
                 f'\\node[anchor=east, font=\\sffamily\\scriptsize] at (-0.05,{py(q):.2f}) {{{q}}};')
    for a in range(2000, 2025, 5):
        o.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize] at ({px(a):.2f},-0.05) {{{a}}};')
    o.append(f'  \\draw[black!70, line width=0.5pt] (0,0) -- ({W:.2f},0);')
    o.append(f'  \\node[rotate=90, font=\\sffamily\\scriptsize] at (-0.65,{H / 2:.2f}) {{\\% de las exportaciones}};')
    for a, t in HITOS:
        o.append(f'  \\draw[black!45, densely dashed, line width=0.4pt] ({px(a):.2f},0) -- ({px(a):.2f},{H:.2f}); '
                 f'\\node[anchor=south, font=\\sffamily\\tiny, text=black!70] at ({px(a):.2f},{H:.2f}) {{{t} ({a})}};')
    camino = ' -- '.join(f'({px(a):.3f},{py(q):.3f})' for a, q in puntos)
    o.append(f'  \\draw[cafr, line width=1.2pt] {camino};')
    for a, q in puntos:
        o.append(f'  \\fill[cafr] ({px(a):.3f},{py(q):.3f}) circle (0.045);')
    # (b) Barras 2023
    X0, Wb, paso = W + 3.2, 5.6, 0.48
    vals = []
    for g, nombre in BLOQUES:
        if g == 'EFTA':
            vals.append((nombre, EFTA_2023, g))
        else:
            vals.append((nombre, serie(wb[f'{g}_intrarregional'])[2023][0], g))
    o.append(f'  \\node[anchor=west, font=\\sffamily\\small] at ({X0 - 1.9:.2f},{H + 0.75:.2f}) {{(b) Comercio intrarregional, 2023}};')
    ytop = H - 0.1
    for t in range(0, 61, 20):
        xx = X0 + t / 60 * Wb
        yb = ytop - (len(vals) - 0.5) * paso
        o.append(f'  \\draw[black!12, line width=0.3pt] ({xx:.2f},{ytop + 0.2:.2f}) -- ({xx:.2f},{yb:.2f}); '
                 f'\\node[anchor=north, font=\\sffamily\\scriptsize] at ({xx:.2f},{yb:.2f}) {{{t}}};')
    for k, (nombre, v, g) in enumerate(vals):
        y = ytop - k * paso
        color = 'cafr' if g == 'AFCFTA' else 'cotro'
        o.append(f'  \\node[anchor=east, font=\\sffamily\\scriptsize] at ({X0 - 0.08:.2f},{y:.2f}) {{{nombre}}};')
        o.append(f'  \\fill[{color}] ({X0:.2f},{y - 0.15:.2f}) rectangle ({X0 + v / 60 * Wb:.3f},{y + 0.15:.2f});')
        cifra = f'{v:.1f}'.replace('.', '{,}')   # coma decimal
        o.append(f'  \\node[anchor=west, font=\\sffamily\\tiny] at ({X0 + v / 60 * Wb + 0.05:.3f},{y:.2f}) {{{cifra}}};')
    o.append('\\end{tikzpicture}')
    texto = '\n'.join(o) + '\n'
    destino = TEMA / 'Img' / 'AfCFTA_comercio_intrarregional.tikz'
    destino.write_text(texto, encoding='utf-8')
    print('Escrito', destino)
    print(puntos[0], puntos[-1], [(n, round(v, 1)) for n, v, _ in vals])


if __name__ == '__main__':
    main()
