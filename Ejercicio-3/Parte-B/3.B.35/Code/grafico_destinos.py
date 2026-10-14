#!/usr/bin/env python3
"""Genera Img/<GRUPO>_destinos_<AÑO>.tikz: barras apiladas (TikZ puro) con el destino de las exportaciones de cada miembro.

Lee la hoja <GRUPO>_destinos_<AÑO> de Code/datos_3B35.xlsx (la escribe Code/datos_comercio.py <GRUPO> destinos <AÑO>).
Uso, desde la carpeta del tema (con openpyxl):
  /tmp/v/bin/python Code/grafico_destinos.py CAN 2023
  bash Code/generar_figura.sh CAN_destinos_2023
"""
import sys
from pathlib import Path

NOMBRES = {'ISL': 'Islandia', 'NOR': 'Noruega', 'CHE': 'Suiza', 'BOL': 'Bolivia', 'COL': 'Colombia', 'ECU': 'Ecuador', 'PER': 'Perú', 'GTM': 'Guatemala',
           'SLV': 'El Salvador', 'HND': 'Honduras', 'NIC': 'Nicaragua', 'CRI': 'Costa Rica'}
SERIES = [  # columna de la hoja, leyenda, color RGB
    (1, 'Resto del bloque', (233, 160, 40)),
    (2, 'EE.\\,UU.', (30, 75, 135)),
    (3, 'China', (200, 60, 60)),
    (4, 'UE-27', (90, 150, 90)),
    (5, 'Resto del mundo', (190, 190, 190)),
]


def main(grupo, anio):
    from openpyxl import load_workbook
    tema = Path(__file__).resolve().parent.parent
    ws = load_workbook(tema / 'Code' / 'datos_3B35.xlsx')[f'{grupo}_destinos_{anio}']
    filas = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    paises = [NOMBRES.get(r[0], r[0]) for r in filas]
    out = [f'% Generado por Code/grafico_destinos.py ({grupo}, {anio}) con la hoja {grupo}_destinos_{anio}.']
    for i, (_, _, rgb) in enumerate(SERIES):
        out.append('\\definecolor{cd%d}{RGB}{%d,%d,%d}' % (i, *rgb))
    ancho, alto_barra, paso, x0 = 13.5, 0.46, 0.72, 2.0   # cm; x0 deja sitio a los nombres
    n = len(filas)
    out.append('\\begin{tikzpicture}[font=\\sffamily]')
    # Leyenda
    lx = x0
    for i, (_, ley, _) in enumerate(SERIES):
        out.append(f'  \\fill[cd{i}] ({lx:.2f},{0.55:.2f}) rectangle ++(0.3,0.2); '
                   f'\\node[anchor=west, font=\\sffamily\\scriptsize, inner sep=1pt] at ({lx + 0.32:.2f},0.65) {{{ley}}};')
        lx += 0.85 + 0.13 * len(ley.replace('\\,', ''))
    # Cuadrícula y eje
    yb = -(n - 1) * paso - alto_barra / 2 - 0.15
    for t in range(0, 101, 10):
        xx = x0 + ancho * t / 100
        out.append(f'  \\draw[black!12, line width=0.3pt] ({xx:.2f},0.25) -- ({xx:.2f},{yb:.2f});')
        out.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize, inner sep=2pt] at ({xx:.2f},{yb:.2f}) {{{t}}};')
    out.append(f'  \\draw[black!70, line width=0.5pt] ({x0:.2f},{yb:.2f}) -- ({x0 + ancho:.2f},{yb:.2f});')
    out.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize] at ({x0 + ancho / 2:.2f},{yb - 0.35:.2f}) '
               '{\\% de las exportaciones de mercancías};')
    # Barras y cifras
    for k, r in enumerate(filas):
        y = -k * paso
        out.append(f'  \\node[anchor=east, font=\\sffamily\\scriptsize] at ({x0 - 0.1:.2f},{y:.2f}) {{{NOMBRES.get(r[0], r[0])}}};')
        acum = 0.0
        for i, (col, _, _) in enumerate(SERIES):
            v = r[col] or 0.0
            xa, xb = x0 + ancho * acum / 100, x0 + ancho * (acum + v) / 100
            out.append(f'  \\fill[cd{i}, draw=white, line width=0.3pt] ({xa:.3f},{y - alto_barra / 2:.3f}) rectangle ({xb:.3f},{y + alto_barra / 2:.3f});')
            if v >= 5:
                color = 'black!80' if i == len(SERIES) - 1 else 'white'
                out.append(f'  \\node[font=\\sffamily\\tiny, text={color}] at ({(xa + xb) / 2:.3f},{y:.2f}) {{{v:.0f}}};')
            acum += v
    out.append('\\end{tikzpicture}')
    destino = tema / 'Img' / f'{grupo}_destinos_{anio}.tikz'
    destino.write_text('\n'.join(out) + '\n', encoding='utf-8')
    print('Escrito', destino)


if __name__ == '__main__':
    main(sys.argv[1], int(sys.argv[2]))
