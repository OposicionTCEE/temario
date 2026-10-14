#!/usr/bin/env python3
"""Genera Img/Megaacuerdos_peso.tikz (TikZ puro) y la hoja Megaacuerdos_2023 de Code/datos_3B35.xlsx: peso del RCEP, el
CPTPP, la UE-27, el USMCA y la ASEAN en el PIB, la población y las exportaciones de bienes y servicios mundiales, 2023.

Fuente: Banco Mundial, Indicadores del Desarrollo Mundial (API https://api.worldbank.org/v2): NY.GDP.MKTP.CD (PIB en
dólares corrientes), SP.POP.TOTL (población) y NE.EXP.GNFS.CD (exportaciones de bienes y servicios). Los datos brutos se
guardan en Code/datos/wdi_2023.csv. Si a un país le falta un dato, no suma en ese indicador (se indica en la hoja).
Uso, desde la carpeta del tema (con openpyxl):
  /tmp/v/bin/python Code/grafico_megaacuerdos.py
  bash Code/generar_figura.sh Megaacuerdos_peso
"""
import csv
import json
import sys
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import datos_comercio as dc  # noqa: E402

TEMA = Path(__file__).resolve().parent.parent
ANIO = 2023
INDICADORES = [('NY.GDP.MKTP.CD', 'PIB'), ('SP.POP.TOTL', 'Población'), ('NE.EXP.GNFS.CD', 'Exportaciones')]
ACUERDOS = [
    ('RCEP', ['BRN', 'KHM', 'IDN', 'LAO', 'MYS', 'MMR', 'PHL', 'SGP', 'THA', 'VNM', 'CHN', 'JPN', 'KOR', 'AUS', 'NZL']),
    ('CPTPP', ['AUS', 'BRN', 'CAN', 'CHL', 'JPN', 'MYS', 'MEX', 'NZL', 'PER', 'SGP', 'VNM', 'GBR']),
    ('UE-27', dc.UE27),
    ('USMCA', ['USA', 'CAN', 'MEX']),
    ('ASEAN', ['BRN', 'KHM', 'IDN', 'LAO', 'MYS', 'MMR', 'PHL', 'SGP', 'THA', 'VNM', 'TLS']),
]
COLORES = [(30, 75, 135), (233, 160, 40), (90, 150, 90)]


def descarga():
    ruta = TEMA / 'Code' / 'datos' / 'wdi_2023.csv'
    datos = {}
    if ruta.exists():
        with open(ruta, newline='', encoding='utf-8') as f:
            for r in csv.DictReader(f):
                datos[(r['pais'], r['indicador'])] = float(r['valor'])
        return datos
    for ind, _ in INDICADORES:
        url = f'https://api.worldbank.org/v2/country/all/indicator/{ind}?date={ANIO}&format=json&per_page=400'
        with urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'}), timeout=90) as r:
            filas = json.loads(r.read().decode('utf-8'))[1]
        for x in filas:
            if x['value'] is not None and x['countryiso3code']:
                datos[(x['countryiso3code'], ind)] = float(x['value'])
    with open(ruta, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['pais', 'indicador', 'valor', 'fuente'])
        for (p, i), v in sorted(datos.items()):
            w.writerow([p, i, v, f'Banco Mundial, WDI ({ANIO})'])
    return datos


def main():
    from openpyxl import load_workbook
    datos = descarga()
    res = []
    for nombre, miembros in ACUERDOS:
        fila, faltan = [], []
        for ind, _ in INDICADORES:
            tot = sum(datos.get((p, ind), 0.0) for p in miembros)
            faltan += [f'{p} ({ind})' for p in miembros if (p, ind) not in datos]
            fila.append(100 * tot / datos[('WLD', ind)])
        res.append((nombre, fila, faltan))
    wb = load_workbook(dc.XLSX)
    if 'Megaacuerdos_2023' in wb.sheetnames:
        del wb['Megaacuerdos_2023']
    ws = wb.create_sheet('Megaacuerdos_2023')
    ws.append([f'Peso de los grandes acuerdos en el mundo, {ANIO} (% del total mundial)'])
    ws.append(['Fuente: Banco Mundial, Indicadores del Desarrollo Mundial (NY.GDP.MKTP.CD, SP.POP.TOTL, NE.EXP.GNFS.CD). '
               'Cálculo propio (Code/grafico_megaacuerdos.py); datos brutos en Code/datos/wdi_2023.csv.'])
    ws.append([])
    ws.append(['Acuerdo'] + [n for _, n in INDICADORES] + ['Datos que faltan'])
    for nombre, fila, faltan in res:
        ws.append([nombre] + [round(v, 2) for v in fila] + [', '.join(faltan)])
    wb.save(dc.XLSX)

    o = ['% Generado por Code/grafico_megaacuerdos.py.']
    for i, rgb in enumerate(COLORES):
        o.append('\\definecolor{cm%d}{RGB}{%d,%d,%d}' % (i, *rgb))
    o.append('\\begin{tikzpicture}[font=\\sffamily]')
    ancho, x0, alto, paso = 11.0, 1.8, 0.22, 1.0
    vmax = 35
    for i, (_, n) in enumerate(INDICADORES):
        o.append(f'  \\fill[cm{i}] ({x0 + i * 2.6:.2f},0.62) rectangle ++(0.3,0.16); '
                 f'\\node[anchor=west, font=\\sffamily\\scriptsize] at ({x0 + i * 2.6 + 0.33:.2f},0.70) {{{n}}};')
    yb = -(len(res) - 1) * paso - 0.55
    for t in range(0, vmax + 1, 5):
        xx = x0 + ancho * t / vmax
        o.append(f'  \\draw[black!12, line width=0.3pt] ({xx:.2f},0.35) -- ({xx:.2f},{yb:.2f}); '
                 f'\\node[anchor=north, font=\\sffamily\\scriptsize] at ({xx:.2f},{yb:.2f}) {{{t}}};')
    o.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize] at ({x0 + ancho / 2:.2f},{yb - 0.35:.2f}) {{\\% del total mundial}};')
    for k, (nombre, fila, _) in enumerate(res):
        y = -k * paso
        o.append(f'  \\node[anchor=east, font=\\sffamily\\scriptsize\\bfseries] at ({x0 - 0.1:.2f},{y:.2f}) {{{nombre}}};')
        for i, v in enumerate(fila):
            yy = y + (1 - i) * (alto + 0.04)
            o.append(f'  \\fill[cm{i}] ({x0:.2f},{yy - alto / 2:.3f}) rectangle ({x0 + ancho * v / vmax:.3f},{yy + alto / 2:.3f});')
            cifra = f'{v:.1f}'.replace('.', '{,}')
            o.append(f'  \\node[anchor=west, font=\\sffamily\\tiny] at ({x0 + ancho * v / vmax + 0.05:.3f},{yy:.3f}) {{{cifra}}};')
    o.append('\\end{tikzpicture}')
    (TEMA / 'Img' / 'Megaacuerdos_peso.tikz').write_text('\n'.join(o) + '\n', encoding='utf-8')
    print('Escrito Img/Megaacuerdos_peso.tikz')
    for nombre, fila, faltan in res:
        print(nombre, [round(v, 1) for v in fila], faltan)


if __name__ == '__main__':
    main()
