#!/usr/bin/env python3
"""Genera Img/CCG_destinos.tikz (TikZ puro) y la hoja CCG_destinos de Code/datos_3B35.xlsx: destino de las exportaciones
del Consejo de Cooperación del Golfo entre 1998 y 2023, según los datos espejo (lo que declaran los importadores).

Los seis miembros declaran la mayor parte de su petróleo sin país de destino, por lo que se usan datos espejo:
Code/datos/espejo.csv (Code/datos_comercio.py PAFTA espejo 1998 2023, que incluye a los seis). Solo cuentan los
importadores que declaran cada año.
Uso, desde la carpeta del tema (con openpyxl):
  /tmp/v/bin/python Code/grafico_ccg.py
  bash Code/generar_figura.sh CCG_destinos
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import datos_comercio as dc  # noqa: E402
from grafico_pafta import lee_espejo  # noqa: E402

TEMA = Path(__file__).resolve().parent.parent
CCG = ['ARE', 'BHR', 'KWT', 'OMN', 'QAT', 'SAU']
ANIOS = [1998, 2003, 2008, 2013, 2018, 2023]
DESTINOS = [  # (leyenda, importadores o None = resto, color)
    ('Resto del CCG', 'CCG', (60, 140, 90)),
    ('Otros países árabes', 'ARABES', (150, 200, 160)),
    ('UE-27', dc.UE27, (130, 130, 200)),
    ('EE.\\,UU.', ['USA'], (30, 75, 135)),
    ('Japón y Corea', ['JPN', 'KOR'], (110, 165, 215)),
    ('China', ['CHN'], (200, 60, 60)),
    ('India', ['IND'], (233, 160, 40)),
    ('Resto del mundo', None, (190, 190, 190)),
]


def reparto(esp, anio):
    arabes = [p for p in dc.GRUPOS['PAFTA'][0] if p not in CCG]
    tot = sum(esp.get((e, anio, 'WLD'), 0.0) for e in CCG)
    fila, usado = [], 0.0
    for _, imp, _ in DESTINOS:
        if imp is None:
            v = tot - usado
        else:
            lista = CCG if imp == 'CCG' else arabes if imp == 'ARABES' else imp
            v = sum(esp.get((e, anio, i), 0.0) for e in CCG for i in lista if i != e)
            usado += v
        fila.append(100 * v / tot)
    return fila, tot


def main():
    from openpyxl import load_workbook
    esp = lee_espejo()
    filas = [(a, *reparto(esp, a)) for a in ANIOS]
    wb = load_workbook(dc.XLSX)
    if 'CCG_destinos' in wb.sheetnames:
        del wb['CCG_destinos']
    ws = wb.create_sheet('CCG_destinos')
    ws.append(['Destino de las exportaciones de mercancías del CCG (% del total), según los importadores (datos espejo)'])
    ws.append([f'Fuente: {dc.FUENTE_ESPEJO}; datos brutos en Code/datos/espejo.csv. Cálculo propio (Code/grafico_ccg.py).'])
    ws.append([])
    ws.append(['Año'] + [d[0].replace('\\,', ' ') for d in DESTINOS] + ['Exportaciones (millones USD)'])
    for a, fila, tot in filas:
        ws.append([a] + [round(v, 2) for v in fila] + [round(tot / 1e6, 1)])
    wb.save(dc.XLSX)

    o = ['% Generado por Code/grafico_ccg.py.']
    for i, (_, _, rgb) in enumerate(DESTINOS):
        o.append('\\definecolor{cd%d}{RGB}{%d,%d,%d}' % (i, *rgb))
    o.append('\\begin{tikzpicture}[font=\\sffamily]')
    ancho, alto, paso, x0 = 13.0, 0.5, 0.75, 1.6
    lx = x0
    for i, (ley, _, _) in enumerate(DESTINOS):
        col, fl = i % 4, i // 4
        xx, yy = x0 + col * 3.3, 1.05 - fl * 0.35
        o.append(f'  \\fill[cd{i}] ({xx:.2f},{yy - 0.08:.2f}) rectangle ++(0.3,0.16); '
                 f'\\node[anchor=west, font=\\sffamily\\scriptsize] at ({xx + 0.33:.2f},{yy:.2f}) {{{ley}}};')
    for k, (a, fila, tot) in enumerate(filas):
        y = -k * paso
        o.append(f'  \\node[anchor=east, font=\\sffamily\\scriptsize] at ({x0 - 0.1:.2f},{y:.2f}) {{{a}}};')
        acum = 0.0
        for i, v in enumerate(fila):
            xa, xb = x0 + ancho * acum / 100, x0 + ancho * (acum + v) / 100
            o.append(f'  \\fill[cd{i}, draw=white, line width=0.3pt] ({xa:.3f},{y - alto / 2:.3f}) rectangle ({xb:.3f},{y + alto / 2:.3f});')
            if v >= 4.5:
                tc = 'black!80' if i in (1, 7) else 'white'
                o.append(f'  \\node[font=\\sffamily\\tiny, text={tc}] at ({(xa + xb) / 2:.3f},{y:.2f}) {{{v:.0f}}};')
            acum += v
    yb = -(len(filas) - 1) * paso - alto / 2 - 0.1
    for t in range(0, 101, 10):
        xx = x0 + ancho * t / 100
        o.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize] at ({xx:.2f},{yb:.2f}) {{{t}}};')
    o.append(f'  \\draw[black!70, line width=0.5pt] ({x0:.2f},{yb:.2f}) -- ({x0 + ancho:.2f},{yb:.2f});')
    o.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize] at ({x0 + ancho / 2:.2f},{yb - 0.35:.2f}) '
             '{\\% de las exportaciones de mercancías, según los países importadores};')
    o.append('\\end{tikzpicture}')
    (TEMA / 'Img' / 'CCG_destinos.tikz').write_text('\n'.join(o) + '\n', encoding='utf-8')
    print('Escrito Img/CCG_destinos.tikz')
    for a, fila, tot in filas:
        print(a, [round(v, 1) for v in fila], round(tot / 1e9, 1))


if __name__ == '__main__':
    main()
