#!/usr/bin/env python3
"""Indicadores económicos de un bloque (3.B.35): intensidad, composición, complementariedad y saldos.

Datos: Code/datos/exportaciones_bilaterales.csv (datos_comercio.py), composicion_cuci.csv, productos_cuci3.csv y
wdi_series.csv (datos_producto.py). Escribe una hoja por indicador en Code/datos_3B35.xlsx, con su fuente.

  - Intensidad del comercio intrarregional: cuota intrarregional / peso del bloque en las importaciones mundiales de
    mercancías (Banco Mundial, WDI). Vale 1 si los socios comercian entre sí lo que predice su peso en el comercio
    mundial. Corrige el tamaño, no la distancia: sirve para seguir la evolución de un bloque, no para comparar niveles
    entre bloques.
  - Composición: exportaciones dentro y fuera del bloque por secciones de la CUCI; manufacturas = secciones 5 a 8.
  - Complementariedad comercial (MICHAELY, 1996): C_ij = 100 (1 - suma_k |m_jk - x_ik| / 2), con x_ik la parte del grupo
    k (3 dígitos de la CUCI) en las exportaciones de i y m_jk su parte en las importaciones de j. 0: nada de lo que
    vende i lo compra j; 100: las dos estructuras coinciden.
  - Saldo con el bloque: exportaciones a los socios menos importaciones desde ellos, medidas con las exportaciones que
    declaran los socios (dato espejo, FOB), en % del PIB (WDI).
Uso, desde la carpeta del tema:  python3 Code/indicadores_bloque.py MCCA
"""
import csv
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import datos_comercio as dc  # noqa: E402
import datos_producto as dp  # noqa: E402

MANUF = {'5', '6', '7', '8'}   # la sección 6 va sin la división 68 (metales no ferrosos), como en la UNCTAD
SECS = list('0123456789') + ['68']
SECCIONES = [  # (secciones de la CUCI, nombre); '68' = metales no ferrosos, restados de la sección 6
    (('0', '1'), 'Alimentos y bebidas'), (('2', '4'), 'Materias primas y aceites'), (('3',), 'Combustibles'),
    (('68',), 'Metales no ferrosos'),
    (('5',), 'Productos químicos'), (('6',), 'Manufacturas por material'), (('7',), 'Maquinaria y transporte'),
    (('8',), 'Manufacturas diversas'), (('9',), 'Oro y otros')]


def wdi():
    out = {}
    for r in dp.lee(dp.CSV_WDI):
        out[(r['pais'], r['indicador'], int(r['anio']))] = float(r['valor'])
    return out


def intensidad(grupo):
    """[(año, cuota %, peso del bloque en las importaciones mundiales %, intensidad)] con los años completos."""
    filas = dc.lee_csv()
    w = wdi()
    anios = sorted({a for (e, a, s) in filas if e in dc.GRUPOS[grupo][0]})
    out = []
    for a, intra, total, cuota, _, _ in dc.cuadro(grupo, anios[0], anios[-1], filas):
        if cuota is None:
            continue
        mm = dc.miembros(grupo, a)
        imp = [w.get((e, 'TM.VAL.MRCH.CD.WT', a)) for e in mm]
        mundo = w.get(('WLD', 'TM.VAL.MRCH.CD.WT', a))
        if None in imp or not mundo:
            continue
        peso = 100 * sum(imp) / mundo
        out.append((a, cuota, peso, cuota / peso))
    return out


def valores():
    """({(exportador, año, socio, sección): valor}, {(exportador, año)} con la división 68 descargada).
    La división 68 (metales no ferrosos) se resta de la sección 6 y va aparte, con la clave '68'."""
    v = defaultdict(float)
    for r in dp.lee(dp.CSV_COMP):
        v[(r['exportador'], int(r['anio']), r['socio'], r['seccion'])] += float(r['valor_usd'])
    con68 = set()
    for r in dp.lee(dp.CSV_68):
        k = (r['exportador'], int(r['anio']), r['socio'])
        x = float(r['valor_usd'])
        con68.add(k[:2])
        v[k + ('68',)] += x
        v[k + ('6',)] = max(v[k + ('6',)] - x, 0.0)
    return v, con68


def completo(v, con68, mm, a):
    """Declaran todos los miembros y se tiene su división 68."""
    return all(any(v.get((e, a, 'WLD', s)) for s in '0123456789') and (e, a) in con68 for e in mm)


def composicion(grupo):
    """{año: (dentro {sección: valor}, fuera {sección: valor})} con los años en que declaran todos los miembros."""
    v, con68 = valores()
    out = {}
    for a in sorted({k[1] for k in v}):
        mm = dc.miembros(grupo, a)
        if not completo(v, con68, mm, a):
            continue
        dentro, fuera = {}, {}
        for s in SECS:
            d = sum(v.get((e, a, p, s), 0.0) for e in mm for p in mm if p != e)
            t = sum(v.get((e, a, 'WLD', s), 0.0) for e in mm)
            dentro[s], fuera[s] = d, max(t - d, 0.0)
        out[a] = (dentro, fuera)
    return out


def reparto_manufacturas(grupo):
    """{año: {miembro: % de las exportaciones de manufacturas dentro del bloque}} (años en que declaran todos)."""
    v, con68 = valores()
    out = {}
    for a in sorted({k[1] for k in v}):
        mm = dc.miembros(grupo, a)
        if not completo(v, con68, mm, a):
            continue
        x = {e: sum(v.get((e, a, p, s), 0.0) for p in mm if p != e for s in MANUF) for e in mm}
        if sum(x.values()):
            out[a] = {e: 100 * x[e] / sum(x.values()) for e in mm}
    return out


def cuota_manufacturas(comp):
    return [(a, 100 * sum(d[s] for s in MANUF) / sum(d.values()), 100 * sum(f[s] for s in MANUF) / sum(f.values()))
            for a, (d, f) in sorted(comp.items()) if sum(d.values()) and sum(f.values())]


def estructuras():
    """{(país, año, flujo): {grupo CUCI de 3 dígitos: parte}}."""
    v = defaultdict(float)
    for r in dp.lee(dp.CSV_PROD):
        v[(r['pais'], int(r['anio']), r['flujo'], r['producto'])] += float(r['valor_usd'])
    tot = defaultdict(float)
    for (p, a, f, k), x in v.items():
        tot[(p, a, f)] += x
    n = defaultdict(int)
    for (p, a, f, k) in v:
        n[(p, a, f)] += 1
    out = defaultdict(dict)
    for (p, a, f, k), x in v.items():
        if tot[(p, a, f)] and n[(p, a, f)] >= 100:   # menos de 100 grupos: declaración incompleta (China, 1985)
            out[(p, a, f)][k] = x / tot[(p, a, f)]
    return out


def michaely(x, m):
    return 100 * (1 - sum(abs(m.get(k, 0.0) - x.get(k, 0.0)) for k in set(x) | set(m)) / 2)


def complementariedad(grupo, referencias=('USA',)):
    """{año: {(exportador, importador): índice}} entre los miembros de cada año y con los mercados de referencia."""
    est = estructuras()
    out = {}
    for a in sorted({k[1] for k in est}):
        mm = dc.miembros(grupo, a)
        c = {}
        for i in mm:
            if (i, a, 'X') not in est:
                continue
            for j in mm + list(referencias):
                if j != i and (j, a, 'M') in est:
                    c[(i, j)] = michaely(est[(i, a, 'X')], est[(j, a, 'M')])
        if c:
            out[a] = c
    return out


def saldos(grupo):
    """[(año, {miembro: saldo con el resto del bloque en % del PIB})] con los años en que declaran todos."""
    filas = dc.lee_csv()
    w = wdi()
    anios = sorted({a for (e, a, s) in filas if e in dc.GRUPOS[grupo][0]})
    out = []
    for a in anios:
        mm = dc.miembros(grupo, a)
        if not all(filas.get((e, a, 'WLD'), (0, ''))[0] for e in mm):
            continue
        res = {}
        for i in mm:
            x = sum(filas.get((i, a, s), (0.0, ''))[0] for s in mm if s != i)
            m = sum(filas.get((s, a, i), (0.0, ''))[0] for s in mm if s != i)
            pib = w.get((i, 'NY.GDP.MKTP.CD', a))
            if pib:
                res[i] = 100 * (x - m) / pib
        if len(res) == len(mm):
            out.append((a, res))
    return out


def hoja(nombre, titulo, fuente, cabecera, filas):
    from openpyxl import Workbook, load_workbook
    wb = load_workbook(dc.XLSX) if dc.XLSX.exists() else Workbook()
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)
    for f in ([titulo], [fuente], [], cabecera, *filas):
        ws.append(list(f))
    wb.save(dc.XLSX)


def main(grupo):
    desc = dc.GRUPOS[grupo][1]
    it = intensidad(grupo)
    comp = composicion(grupo)
    cm = cuota_manufacturas(comp)
    cp = complementariedad(grupo)
    sd = saldos(grupo)
    rm = reparto_manufacturas(grupo)
    try:
        hoja(f'{grupo}_intensidad', f'Intensidad del comercio intrarregional. {desc}',
             'Cuota: exportaciones a los socios / exportaciones totales (Code/datos_comercio.py). Peso: importaciones de '
             'mercancías de los miembros / importaciones mundiales (Banco Mundial, WDI, TM.VAL.MRCH.CD.WT). '
             'Intensidad = cuota / peso. Cálculo propio (Code/indicadores_bloque.py).',
             ['Año', 'Cuota intrarregional (%)', 'Peso en las importaciones mundiales (%)', 'Intensidad'],
             [(a, round(c, 2), round(p, 4), round(i, 1)) for a, c, p, i in it])
        hoja(f'{grupo}_composicion', f'Exportaciones dentro y fuera del bloque por secciones de la CUCI. {desc}',
             'UN Comtrade, API pública: CUCI Rev. 1 (hasta 1995), Rev. 3 (1996-2006) y Rev. 4 (desde 2007); exportaciones '
             'declaradas por cada miembro al mundo y a cada socio, dólares corrientes, en millones. Fuera = mundo menos socios. '
             'La división 68 (metales no ferrosos) se resta de la sección 6, como en la definición de manufacturas de la UNCTAD. '
             'Cálculo propio (Code/indicadores_bloque.py).',
             ['Año', 'Destino'] + [f'Sección {s}' + (' (sin la división 68)' if s == '6' else '') for s in '0123456789']
             + ['División 68 (metales no ferrosos)', 'Manufacturas (5-8 sin 68, %)'],
             [(a, dest, *[round(x[s] / 1e6, 2) for s in SECS], round(100 * sum(x[s] for s in MANUF) / sum(x.values()), 1))
              for a, (d, f) in sorted(comp.items()) for dest, x in (('Dentro', d), ('Fuera', f)) if sum(x.values())])
        hoja(f'{grupo}_complementariedad', f'Complementariedad comercial (MICHAELY, 1996). {desc}',
             'UN Comtrade, API pública, grupos de 3 dígitos de la CUCI (Rev. 1 hasta 1995, Rev. 3 de 1996 a 2006 y Rev. 4 '
             'desde 2007): exportaciones del primer país y importaciones del segundo, con el mundo. Cálculo propio '
             '(Code/indicadores_bloque.py).',
             ['Año', 'Exportador', 'Importador', 'Índice (0-100)'],
             [(a, i, j, round(x, 1)) for a, c in sorted(cp.items()) for (i, j), x in sorted(c.items())])
        hoja(f'{grupo}_saldos', f'Saldo comercial de cada miembro con el resto del bloque, % del PIB. {desc}',
             'Exportaciones a los socios menos importaciones desde ellos (exportaciones declaradas por los socios), '
             'con Code/datos_comercio.py; PIB en dólares corrientes del Banco Mundial (WDI). Cálculo propio.',
             ['Año'] + dc.GRUPOS[grupo][0], [(a, *[round(r.get(e), 2) if e in r else None for e in dc.GRUPOS[grupo][0]]) for a, r in sd])
        hoja(f'{grupo}_reparto_manuf', f'Reparto entre los miembros de las exportaciones de manufacturas dentro del bloque (%). {desc}',
             'UN Comtrade, API pública (CUCI, secciones 5 a 8; revisiones como en la hoja de composición). Cálculo propio.',
             ['Año'] + dc.GRUPOS[grupo][0], [(a, *[round(r[e], 1) if e in r else None for e in dc.GRUPOS[grupo][0]]) for a, r in sorted(rm.items())])
    except ImportError:
        print('Falta openpyxl: no se escribe el Excel')
    return it, comp, cm, cp, sd, rm


if __name__ == '__main__':
    it, comp, cm, cp, sd, rm = main(sys.argv[1])
    print('Reparto manufacturas:', [(a, {k: round(v) for k, v in r.items()}) for a, r in sorted(rm.items())][::5])
    print('Intensidad:', [(a, round(c, 1), round(i)) for a, c, p, i in it][::5])
    print('Manufacturas (dentro, fuera):', [(a, round(d), round(f)) for a, d, f in cm][::4])
    for a, c in cp.items():
        intra = [x for (i, j), x in c.items() if j in dc.GRUPOS[sys.argv[1]][0]]
        ref = [x for (i, j), x in c.items() if j not in dc.GRUPOS[sys.argv[1]][0]]
        print('Complementariedad', a, round(sum(intra) / len(intra), 1) if intra else None, round(sum(ref) / len(ref), 1) if ref else None)
    print('Saldos:', [(a, {k: round(v, 1) for k, v in r.items()}) for a, r in sd][::6])
