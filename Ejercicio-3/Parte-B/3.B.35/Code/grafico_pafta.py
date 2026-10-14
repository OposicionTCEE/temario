#!/usr/bin/env python3
"""Genera Img/PAFTA_comercio.tikz (TikZ puro, dos paneles) y dos hojas de Code/datos_3B35.xlsx:
(a) cuota de las exportaciones a los demás miembros de la PAFTA en las exportaciones totales, 1998-2023, para todos los
    países que declaran y sin los seis del Consejo de Cooperación del Golfo (hoja PAFTA_con_y_sin_CCG);
(b) destino de las exportaciones de cada subregión en 2023 (hoja PAFTA_subregiones_2023).

Datos: Code/datos/exportaciones_bilaterales.csv y Code/datos/destinos.csv (los escribe Code/datos_comercio.py:
  PAFTA destinos 2023 y PAFTA 1998 2023). Solo cuentan los países que declaran cada año; se omiten los años en que
  declara menos de la mitad del grupo correspondiente.
Uso, desde la carpeta del tema (con openpyxl):
  /tmp/v/bin/python Code/grafico_pafta.py
  bash Code/generar_figura.sh PAFTA_comercio
"""
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import datos_comercio as dc  # noqa: E402

TEMA = Path(__file__).resolve().parent.parent
CCG = ['ARE', 'BHR', 'KWT', 'OMN', 'QAT', 'SAU']
SUBREGIONES = [
    ('Golfo (CCG)', CCG),
    ('Magreb', ['DZA', 'LBY', 'MAR', 'TUN']),
    ('Oriente Próximo, Sudán y Yemen', ['EGY', 'IRQ', 'JOR', 'LBN', 'PSE', 'SYR', 'SDN', 'YEM']),
]
DESTINOS = [  # (leyenda, socios o None = resto, color)
    ('Resto de la PAFTA', 'PAFTA', (233, 160, 40)),
    ('UE-27', dc.UE27, (90, 150, 90)),
    ('China', ['CHN'], (200, 60, 60)),
    ('India', ['IND'], (150, 90, 160)),
    ('Japón y Corea', ['JPN', 'KOR'], (110, 165, 215)),
    ('EE.\\,UU.', ['USA'], (30, 75, 135)),
    ('Resto del mundo', None, (190, 190, 190)),
]


def serie(filas, exportadores):
    out = []
    for anio in range(1998, 2024):
        mm = dc.miembros('PAFTA', anio)
        con = [e for e in exportadores if e in mm and filas.get((e, anio, 'WLD'), (0.0, ''))[0]]
        if len(con) < len([e for e in exportadores if e in mm]) / 2:
            out.append((anio, None, len(con)))
            continue
        tot = sum(filas[(e, anio, 'WLD')][0] for e in con)
        intra = sum(filas.get((e, anio, s), (0.0, ''))[0] for e in con for s in mm if s != e)
        out.append((anio, 100 * intra / tot, len(con)))
    return out


def destinos_subregion(anio=2023):
    brutos = {}
    with open(TEMA / 'Code' / 'datos' / 'destinos.csv', newline='', encoding='utf-8') as f:
        for r in csv.DictReader(f):
            if int(r['anio']) == anio:
                brutos[(r['exportador'], r['socio'])] = float(r['valor_usd'])
    mm = dc.miembros('PAFTA', anio)
    res = []
    for nombre, paises in SUBREGIONES:
        con = [e for e in paises if brutos.get((e, 'WLD'))]
        tot = sum(brutos[(e, 'WLD')] for e in con)
        fila, usado = [], 0.0
        for ley, socios, _ in DESTINOS:
            if socios is None:
                v = tot - usado
            else:
                lista = [s for s in mm] if socios == 'PAFTA' else socios
                v = sum(brutos.get((e, s), 0.0) for e in con for s in lista if s != e)
            usado += v if socios is not None else 0
            fila.append(100 * v / tot)
        res.append((nombre, fila, con, tot))
    return res


def main():
    from openpyxl import load_workbook
    filas = dc.lee_csv()
    todos = serie(filas, dc.GRUPOS['PAFTA'][0])
    sin_ccg = serie(filas, [e for e in dc.GRUPOS['PAFTA'][0] if e not in CCG])
    sub = destinos_subregion()
    wb = load_workbook(dc.XLSX)
    for nombre in ('PAFTA_con_y_sin_CCG', 'PAFTA_subregiones_2023'):
        if nombre in wb.sheetnames:
            del wb[nombre]
    ws = wb.create_sheet('PAFTA_con_y_sin_CCG')
    ws.append(['Exportaciones a los demás miembros de la PAFTA / exportaciones totales (%), países que declaran'])
    ws.append([f'Fuente: {dc.FUENTE_WITS}. Cálculo propio (Code/grafico_pafta.py). Sin cuota si declara menos de la mitad.'])
    ws.append([])
    ws.append(['Año', 'Todos (%)', 'Países que declaran', 'Sin el CCG (%)', 'Países que declaran (sin CCG)'])
    for (a, q1, n1), (_, q2, n2) in zip(todos, sin_ccg):
        ws.append([a, None if q1 is None else round(q1, 2), n1, None if q2 is None else round(q2, 2), n2])
    ws = wb.create_sheet('PAFTA_subregiones_2023')
    ws.append(['Destino de las exportaciones de mercancías de cada subregión de la PAFTA en 2023 (% del total)'])
    ws.append([f'Fuente: {dc.FUENTE_WITS}; datos brutos en Code/datos/destinos.csv. Cálculo propio (Code/grafico_pafta.py).'])
    ws.append([])
    ws.append(['Subregión', 'Países que declaran'] + [d[0].replace('\\,', ' ') for d in DESTINOS] + ['Exportaciones (millones USD)'])
    for nombre, fila, con, tot in sub:
        ws.append([nombre, ', '.join(con)] + [round(v, 2) for v in fila] + [round(tot / 1e6, 1)])
    wb.save(dc.XLSX)

    o = ['% Generado por Code/grafico_pafta.py.', '\\definecolor{ctodos}{RGB}{30,75,135}', '\\definecolor{csin}{RGB}{233,160,40}']
    for i, (_, _, rgb) in enumerate(DESTINOS):
        o.append('\\definecolor{cd%d}{RGB}{%d,%d,%d}' % (i, *rgb))
    o.append('\\begin{tikzpicture}[font=\\sffamily]')
    # (a)
    W, H, a0, a1, ymax = 6.8, 4.0, 1997, 2024, 30

    def px(a):
        return (a - a0) / (a1 - a0) * W

    def py(q):
        return q / ymax * H
    o.append(f'  \\node[anchor=west, font=\\sffamily\\small] at (-0.6,{H + 0.6:.2f}) {{(a) Comercio entre los miembros}};')
    for q in range(0, ymax + 1, 10):
        o.append(f'  \\draw[black!12, line width=0.3pt] (0,{py(q):.2f}) -- ({W:.2f},{py(q):.2f}); '
                 f'\\node[anchor=east, font=\\sffamily\\scriptsize] at (-0.05,{py(q):.2f}) {{{q}}};')
    for a in range(2000, 2025, 5):
        o.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize] at ({px(a):.2f},-0.05) {{{a}}};')
    o.append(f'  \\draw[black!70, line width=0.5pt] (0,0) -- ({W:.2f},0);')
    o.append(f'  \\node[rotate=90, font=\\sffamily\\scriptsize] at (-0.65,{H / 2:.2f}) {{\\% de las exportaciones}};')
    for a, t in [(2005, 'libre comercio pleno (2005)'), (2011, 'Primavera Árabe (2011)')]:
        o.append(f'  \\draw[black!45, densely dashed, line width=0.4pt] ({px(a):.2f},0) -- ({px(a):.2f},{H:.2f}); '
                 f'\\node[anchor=south, font=\\sffamily\\tiny, text=black!70] at ({px(a):.2f},{H:.2f}) {{{t}}};')
    for datos, color, ley, yl in [(todos, 'ctodos', 'todos los miembros', 0.55), (sin_ccg, 'csin', 'sin el Golfo (CCG)', 0.2)]:
        pts = [(a, q) for a, q, _ in datos if q is not None]
        tramos, actual = [], []
        for a, q in pts:   # corta la línea si falta algún año
            if actual and a != actual[-1][0] + 1:
                tramos.append(actual)
                actual = []
            actual.append((a, q))
        if actual:
            tramos.append(actual)
        for tr in tramos:
            if len(tr) > 1:
                o.append(f'  \\draw[{color}, line width=1.2pt] ' + ' -- '.join(f'({px(a):.3f},{py(q):.3f})' for a, q in tr) + ';')
        for a, q in pts:
            o.append(f'  \\fill[{color}] ({px(a):.3f},{py(q):.3f}) circle (0.045);')
        o.append(f'  \\draw[{color}, line width=1.2pt] ({W - 2.3:.2f},{H - yl:.2f}) -- ++(0.35,0); '
                 f'\\node[anchor=west, font=\\sffamily\\tiny] at ({W - 1.9:.2f},{H - yl:.2f}) {{{ley}}};')
    # (b)
    X0, Wb, paso, alto = W + 3.6, 6.0, 0.95, 0.5
    o.append(f'  \\node[anchor=west, font=\\sffamily\\small] at ({X0 - 2.6:.2f},{H + 0.6:.2f}) {{(b) Destino de las exportaciones, 2023}};')
    lx, ly = X0 - 2.5, H - 0.05
    for i, (ley, _, _) in enumerate(DESTINOS):
        col, fila_l = i % 4, i // 4
        xx = lx + col * 2.2
        yy = ly - fila_l * 0.32
        o.append(f'  \\fill[cd{i}] ({xx:.2f},{yy - 0.08:.2f}) rectangle ++(0.28,0.16); '
                 f'\\node[anchor=west, font=\\sffamily\\tiny] at ({xx + 0.3:.2f},{yy:.2f}) {{{ley}}};')
    ytop = H - 1.05
    for k, (nombre, fila, _, _) in enumerate(sub):
        y = ytop - k * paso
        o.append(f'  \\node[anchor=east, align=right, font=\\sffamily\\scriptsize, text width=2.4cm] at ({X0 - 0.08:.2f},{y:.2f}) {{{nombre}}};')
        acum = 0.0
        for i, v in enumerate(fila):
            xa, xb = X0 + Wb * acum / 100, X0 + Wb * (acum + v) / 100
            o.append(f'  \\fill[cd{i}, draw=white, line width=0.3pt] ({xa:.3f},{y - alto / 2:.3f}) rectangle ({xb:.3f},{y + alto / 2:.3f});')
            if v >= 6:
                tc = 'black!80' if i == len(fila) - 1 else 'white'
                o.append(f'  \\node[font=\\sffamily\\tiny, text={tc}] at ({(xa + xb) / 2:.3f},{y:.2f}) {{{v:.0f}}};')
            acum += v
    yb = ytop - (len(sub) - 1) * paso - alto / 2 - 0.1
    for t in range(0, 101, 25):
        xx = X0 + Wb * t / 100
        o.append(f'  \\node[anchor=north, font=\\sffamily\\scriptsize] at ({xx:.2f},{yb:.2f}) {{{t}}};')
    o.append(f'  \\draw[black!70, line width=0.5pt] ({X0:.2f},{yb:.2f}) -- ({X0 + Wb:.2f},{yb:.2f});')
    o.append('\\end{tikzpicture}')
    (TEMA / 'Img' / 'PAFTA_comercio.tikz').write_text('\n'.join(o) + '\n', encoding='utf-8')
    print('Escrito Img/PAFTA_comercio.tikz')
    for (a, q1, n1), (_, q2, n2) in zip(todos, sin_ccg):
        print(a, None if q1 is None else round(q1, 1), n1, None if q2 is None else round(q2, 1), n2)
    for nombre, fila, con, tot in sub:
        print(nombre, [round(v, 1) for v in fila], con, round(tot / 1e9, 1))


if __name__ == '__main__':
    main()
