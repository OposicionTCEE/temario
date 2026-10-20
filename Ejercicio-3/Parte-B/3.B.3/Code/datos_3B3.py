#!/usr/bin/env python3
"""Monta Code/datos_3B3.xlsx, el libro único con los datos que cita el tema 3.B.3 (regla del usuario, 15/10/2026).

Hojas (formato largo, con filtros):
  Léeme           qué hay en cada hoja y cómo se regenera.
  Fuentes         cada fuente con su código: institución, base de datos, enlace.
  Series_brutas   las observaciones descargadas tal cual (código de serie, periodo, valor, unidad).
  Cifras_citadas  cada cifra del texto, con su bloque, la frase donde aparece, el valor redondeado y su fuente.

Descarga las series del ECB Data Portal (API SDMX) y de FRED (Reserva Federal, cuentas Z.1).
Uso, desde la carpeta del tema:  /ruta/a/python-con-openpyxl Code/datos_3B3.py
"""
import csv
import io
import urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font

FECHA_CONSULTA = "2026-10-20"

FUENTES = [
    ("BCE-BSI", "Banco Central Europeo", "Estadísticas de balance de las IFM (BSI)", "https://data.ecb.europa.eu/"),
    ("BCE-QSA", "Banco Central Europeo", "Cuentas financieras trimestrales de los sectores (QSA)", "https://data.ecb.europa.eu/"),
    ("FED-Z1", "Board of Governors of the Federal Reserve System", "Financial Accounts of the United States (Z.1), vía FRED",
     "https://fred.stlouisfed.org/"),
    ("BCE-SAFE", "Banco Central Europeo", "Durante, Ferrando, Grønlund y Reinelt (2023), Economic Bulletin 1/2023, recuadro 7",
     "https://www.ecb.europa.eu/press/economic-bulletin/focus/2023/html/ecb.ebbox202301_07~78a838eb48.en.html"),
]

# (fuente, conjunto, clave, periodo, descripción, unidad)
SERIES = [
    ("BCE-BSI", "BSI", "M.U2.N.A.A20.A.1.U2.2240.Z01.E", "2025-12",
     "Préstamos de las IFM (excl. Eurosistema) a las sociedades no financieras de la zona del euro, saldo", "millones de euros"),
    ("BCE-BSI", "BSI", "M.ES.N.A.A20.A.1.U6.2240.Z01.E", "2025-12",
     "Préstamos de las IFM residentes en España a las sociedades no financieras residentes, saldo", "millones de euros"),
    ("BCE-QSA", "QSA", "Q.N.I9.W0.S11.S1.N.L.LE.F3.T._Z.XDC._T.S.V.N._T", "2025-Q4",
     "Valores de deuda emitidos por las sociedades no financieras de la zona del euro, saldo", "millones de euros"),
    ("BCE-QSA", "QSA", "Q.N.ES.W0.S11.S1.N.L.LE.F3.T._Z.XDC._T.S.V.N._T", "2025-Q4",
     "Valores de deuda emitidos por las sociedades no financieras residentes en España, saldo", "millones de euros"),
]
FRED = [
    ("FED-Z1", "NCBDBIQ027S", "2025-10-01", "Sociedades no financieras de EE. UU.: valores de deuda, pasivo, saldo (4.º trim. 2025)",
     "millones de dólares"),
    ("FED-Z1", "NCBLL", "2025-10-01", "Sociedades no financieras de EE. UU.: préstamos, pasivo, saldo (4.º trim. 2025)",
     "millones de dólares"),
]


def ecb(conjunto, clave, periodo):
    url = (f"https://data-api.ecb.europa.eu/service/data/{conjunto}/{clave}"
           f"?startPeriod={periodo}&endPeriod={periodo}&format=csvdata")
    filas = list(csv.DictReader(io.StringIO(urllib.request.urlopen(url, timeout=60).read().decode())))
    return float(filas[0]["OBS_VALUE"])


def fred(serie, fecha):
    url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={serie}"
    for fila in csv.reader(io.StringIO(urllib.request.urlopen(url, timeout=60).read().decode())):
        if fila and fila[0] == fecha:
            return float(fila[1])
    raise ValueError(f"{serie}: sin dato para {fecha}")


def main():
    brutas = []
    for fuente, conjunto, clave, periodo, desc, unidad in SERIES:
        brutas.append((fuente, f"{conjunto}.{clave}", periodo, ecb(conjunto, clave, periodo), unidad, desc))
    for fuente, serie, fecha, desc, unidad in FRED:
        brutas.append((fuente, serie, fecha, fred(serie, fecha), unidad, desc))
    v = {b[1]: b[3] for b in brutas}

    ea_prest = v["BSI.M.U2.N.A.A20.A.1.U2.2240.Z01.E"]
    es_prest = v["BSI.M.ES.N.A.A20.A.1.U6.2240.Z01.E"]
    ea_val = v["QSA.Q.N.I9.W0.S11.S1.N.L.LE.F3.T._Z.XDC._T.S.V.N._T"]
    es_val = v["QSA.Q.N.ES.W0.S11.S1.N.L.LE.F3.T._Z.XDC._T.S.V.N._T"]
    us_val, us_prest = v["NCBDBIQ027S"], v["NCBLL"]
    frase = "Bloque 1, «Financiación ajena»: préstamos bancarios frente a valores de deuda"
    cifras = [
        ("1", frase, "Zona del euro, préstamos de los bancos a las SNF", round(ea_prest / 1e6, 1), "billones de euros", "BCE-BSI"),
        ("1", frase, "Zona del euro, valores de deuda de las SNF", round(ea_val / 1e6, 1), "billones de euros", "BCE-QSA"),
        ("1", frase, "Zona del euro, cociente préstamos / valores", round(ea_prest / ea_val, 2), "veces", "BCE-BSI, BCE-QSA"),
        ("1", frase, "España, préstamos de los bancos a las SNF", round(es_prest / 1e3), "mil millones de euros", "BCE-BSI"),
        ("1", frase, "España, valores de deuda de las SNF", round(es_val / 1e3), "mil millones de euros", "BCE-QSA"),
        ("1", frase, "España, cociente préstamos / valores", round(es_prest / es_val, 2), "veces", "BCE-BSI, BCE-QSA"),
        ("1", frase, "EE. UU., valores de deuda de las SNF", round(us_val / 1e6, 1), "billones de dólares", "FED-Z1"),
        ("1", frase, "EE. UU., préstamos de las SNF", round(us_prest / 1e6, 1), "billones de dólares", "FED-Z1"),
        ("1", frase, "EE. UU., cociente valores / préstamos", round(us_val / us_prest, 2), "veces", "FED-Z1"),
        ("1", "Bloque 1, «Financiación ajena»: empresas que emiten bonos (2009-2022)", "Empresas de la SAFE que emiten bonos",
         9, "%", "BCE-SAFE"),
        ("1", "Bloque 1, «Financiación ajena»: empresas que usan productos bancarios (2009-2022)",
         "Empresas de la SAFE que usan productos bancarios", 50, "% (algo más del)", "BCE-SAFE"),
        ("Anexo II", "Banca frente a mercado, criterio «Volumen»",
         "Diferencia en la frecuencia de emisión de bonos, empresas de más de 50 empleados frente a las demás", 4,
         "puntos porcentuales", "BCE-SAFE"),
    ]

    wb = Workbook()
    hojas = {
        "Léeme": (["Hoja", "Contenido"], [
            ("Fuentes", "Fuentes con su código."),
            ("Series_brutas", "Observaciones descargadas tal cual, consultadas el " + FECHA_CONSULTA + "."),
            ("Cifras_citadas", "Cada cifra del texto, redondeada como aparece, con su fuente."),
            ("Regenerar", "Ejecutar Code/datos_3B3.py desde la carpeta del tema con un Python que tenga openpyxl."),
        ]),
        "Fuentes": (["Código", "Institución", "Base de datos o documento", "Enlace"], FUENTES),
        "Series_brutas": (["Fuente", "Serie", "Periodo", "Valor", "Unidad", "Descripción"], brutas),
        "Cifras_citadas": (["Bloque", "Dónde", "Indicador", "Valor", "Unidad", "Fuente"], cifras),
    }
    wb.remove(wb.active)
    for nombre, (cab, filas) in hojas.items():
        ws = wb.create_sheet(nombre)
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True)
        for f in filas:
            ws.append(list(f))
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(80, max(12, max(len(str(c.value or "")) for c in col) + 2))
    wb.save("Code/datos_3B3.xlsx")
    for c in cifras:
        print(c[2], "=", c[3], c[4])


if __name__ == "__main__":
    main()
